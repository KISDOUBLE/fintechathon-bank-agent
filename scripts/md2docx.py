# -*- coding: utf-8 -*-
"""极简 Markdown → DOCX 转换器：标题、段落、列表、表格、代码块、行内粗体/代码。"""

import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY = RGBColor(0x12, 0x1A, 0x30)
INK = RGBColor(0x22, 0x28, 0x36)
MUTED = RGBColor(0x6B, 0x76, 0x8C)
CODE_BG = "F3F5F9"
HEADER_BG = "121A30"
BODY_FONT = "宋体"
HEAD_FONT = "微软雅黑"
MONO = "Consolas"


def set_run_font(run, name=BODY_FONT, size=10.5, bold=None, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), name)


def shade(paragraph_or_cell, hex_fill):
    el = paragraph_or_cell._element
    pr = el.find(qn("w:pPr"))
    if pr is None and el.tag.endswith("tc"):
        pr = el.find(qn("w:tcPr"))
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    if el.tag.endswith("tc"):
        el.find(qn("w:tcPr")).append(shd)
    else:
        pr.append(shd)


def shade_cell(cell, hex_fill):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def add_inline(paragraph, text, size=10.5, color=INK, base_bold=False):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            inner = part[2:-2]
            if inner.startswith("`") and inner.endswith("`") and len(inner) > 2:
                run = paragraph.add_run(inner[1:-1])
                set_run_font(run, MONO, size - 0.5, True, color)
            else:
                run = paragraph.add_run(inner)
                set_run_font(run, BODY_FONT, size, True, color)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, MONO, size - 0.5, False, RGBColor(0x2F, 0x5B, 0xFF))
        else:
            run = paragraph.add_run(part)
            set_run_font(run, BODY_FONT, size, base_bold, color)


def build(md_path, docx_path):
    with open(md_path, encoding="utf-8") as f:
        lines = f.read().split("\n")

    doc = Document()
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    style = doc.styles["Normal"]
    style.font.name = BODY_FONT
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    style.paragraph_format.line_spacing = 1.45
    style.paragraph_format.space_after = Pt(4)

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i].rstrip()

        if not line.strip():
            i += 1
            continue

        # 代码块
        if line.startswith("```"):
            i += 1
            buf = []
            while i < n and not lines[i].startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            shade(p, CODE_BG)
            for idx, code_line in enumerate(buf):
                run = p.add_run(code_line)
                set_run_font(run, MONO, 8.5, False, RGBColor(0x1F, 0x27, 0x3A))
                if idx < len(buf) - 1:
                    run.add_break()
            continue

        # 标题
        m = re.match(r"^(#{1,4})\s+(.*)$", line)
        if m:
            level, text = len(m.group(1)), m.group(2).strip()
            p = doc.add_paragraph()
            if level == 1:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(14)
                add_inline(p, text, 20, NAVY, True)
            elif level == 2:
                p.paragraph_format.space_before = Pt(16)
                p.paragraph_format.space_after = Pt(8)
                add_inline(p, text, 15, NAVY, True)
            else:
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(4)
                add_inline(p, text, 12, NAVY, True)
            i += 1
            continue

        # 表格
        if line.strip().startswith("|"):
            block = []
            while i < n and lines[i].strip().startswith("|"):
                block.append(lines[i].strip())
                i += 1
            rows = []
            for row in block:
                cells = [c.strip() for c in row.strip("|").split("|")]
                if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    continue
                rows.append(cells)
            if rows:
                cols = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=cols)
                table.style = "Table Grid"
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                for r, row in enumerate(rows):
                    for c in range(cols):
                        cell = table.cell(r, c)
                        cell.text = ""
                        p = cell.paragraphs[0]
                        p.paragraph_format.space_after = Pt(0)
                        p.paragraph_format.line_spacing = 1.15
                        val = row[c] if c < len(row) else ""
                        add_inline(p, val, 9, RGBColor(0xFF, 0xFF, 0xFF) if r == 0 else INK, r == 0)
                        if r == 0:
                            shade_cell(cell, HEADER_BG)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        # 引用
        if line.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.4)
            add_inline(p, line.lstrip("> ").strip(), 9.5, MUTED)
            i += 1
            continue

        # 分隔线
        if re.fullmatch(r"-{3,}", line.strip()):
            i += 1
            continue

        # 无序列表
        m = re.match(r"^[-*]\s+(.*)$", line)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, m.group(1))
            i += 1
            continue

        # 有序列表
        m = re.match(r"^(\d+)\.\s+(.*)$", line)
        if m:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, m.group(2))
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        add_inline(p, line)
        i += 1

    doc.save(docx_path)
    print("saved", docx_path)


if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2])
